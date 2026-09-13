from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from tests.conftest import register


def create_jijihong_context(client: TestClient) -> tuple[dict[str, str], dict[str, str]]:
    feida_token = register(client, email="brands@example.com", tenant_name="费大厨")
    feida_headers = {"Authorization": f"Bearer {feida_token}"}
    created = client.post(
        "/api/v1/auth/tenants",
        headers=feida_headers,
        json={"name": "季季红", "business_mode": "JIJIHONG"},
    )
    assert created.status_code == 201, created.text
    switched = client.post(
        "/api/v1/auth/switch-tenant",
        headers=feida_headers,
        json={"tenant_id": created.json()["id"]},
    )
    assert switched.status_code == 200, switched.text
    jijihong_headers = {
        "Authorization": f"Bearer {switched.json()['access_token']}"
    }
    return feida_headers, jijihong_headers


def create_device(client: TestClient, headers: dict[str, str], name: str = "手机01") -> int:
    response = client.post(
        "/api/v1/alipay-pool/devices",
        headers=headers,
        json={"name": name},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def create_account(
    client: TestClient,
    headers: dict[str, str],
    *,
    device_id: int,
    alias: str,
    balance: str,
) -> int:
    response = client.post(
        "/api/v1/alipay-pool/accounts",
        headers=headers,
        json={
            "device_id": device_id,
            "alias": alias,
            "current_balance": balance,
            "status": "ACTIVE",
            "phase": "DAY2_ACTIVE",
            "create_coupon": False,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def preview_plan(client: TestClient, headers: dict[str, str], amount: str) -> dict:
    response = client.post(
        "/api/v1/alipay-pool/recommendations/preview",
        headers=headers,
        json={"order_amount": amount},
    )
    assert response.status_code == 200, response.text
    return response.json()


def create_reserved_order(
    client: TestClient,
    headers: dict[str, str],
    *,
    amount: str,
    plan: dict | None = None,
    customer_received: str | None = None,
) -> dict:
    plan = plan or preview_plan(client, headers, amount)
    response = client.post(
        "/api/v1/alipay-pool/orders",
        headers=headers,
        json={
            "business_date": date.today().isoformat(),
            "order_amount": amount,
            "customer_received_amount": customer_received or amount,
            "allocations": [
                {
                    "account_id": item["account_id"],
                    "coupon_amount": item["coupon_amount"],
                    "balance_amount": item["balance_amount"],
                    "external_cash_amount": item["external_cash_amount"],
                }
                for item in plan["allocations"]
            ],
            "external_cash_override_reason": "业务确认" if plan["soft_cap_exceeded"] else None,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_brand_switch_and_pool_isolation(client: TestClient):
    feida_headers, jijihong_headers = create_jijihong_context(client)

    feida_me = client.get("/api/v1/auth/me", headers=feida_headers)
    jijihong_me = client.get("/api/v1/auth/me", headers=jijihong_headers)
    assert feida_me.json()["business_mode"] == "FEDAICHU"
    assert jijihong_me.json()["business_mode"] == "JIJIHONG"
    assert {item["business_mode"] for item in jijihong_me.json()["tenants"]} == {
        "FEDAICHU",
        "JIJIHONG",
    }
    assert client.get(
        "/api/v1/alipay-pool/settings", headers=feida_headers
    ).status_code == 404
    assert client.get(
        "/api/v1/alipay-pool/settings", headers=jijihong_headers
    ).status_code == 200
    assert client.get("/api/v1/partners/sources", headers=jijihong_headers).status_code == 404


def test_split_reserve_confirm_and_reverse(client: TestClient):
    _, headers = create_jijihong_context(client)
    device_id = create_device(client, headers)
    first_id = create_account(
        client, headers, device_id=device_id, alias="acc-70", balance="70"
    )
    second_id = create_account(
        client, headers, device_id=device_id, alias="acc-50", balance="50"
    )

    plan = preview_plan(client, headers, "120")
    assert plan["total_balance"] == "120.00"
    assert plan["total_external_cash"] == "0.00"
    assert {item["account_id"] for item in plan["allocations"]} == {
        first_id,
        second_id,
    }
    preview_accounts = client.get(
        "/api/v1/alipay-pool/accounts", headers=headers
    ).json()
    assert sum(float(item["reserved_balance"]) for item in preview_accounts) == 0
    created = create_reserved_order(client, headers, amount="120", plan=plan)
    order = created["order"]
    reserved_plan = created["payment_plan"]
    assert order["source_id"] is None
    assert order["contractor_id"] is None
    reserved_accounts = client.get("/api/v1/alipay-pool/accounts", headers=headers).json()
    assert sum(float(item["reserved_balance"]) for item in reserved_accounts) == 120
    assert sum(float(item["current_balance"]) for item in reserved_accounts) == 120

    confirmed = client.post(
        f"/api/v1/alipay-pool/payment-plans/{reserved_plan['id']}/confirm",
        headers=headers,
        json={},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "CONFIRMED"
    confirmed_order = client.get(
        f"/api/v1/orders/{order['id']}", headers=headers
    ).json()
    assert confirmed_order["status"] == "SUCCESS"
    assert confirmed_order["actual_paid"] == "120.00"
    assert confirmed_order["customer_received_amount"] == "120.00"
    financial = client.get(
        "/api/v1/alipay-pool/funds/financial-entries", headers=headers
    ).json()["items"]
    assert {item["entry_type"] for item in financial} == {
        "CUSTOMER_RECEIPT",
        "BALANCE_COST",
    }
    after_confirm = client.get(
        "/api/v1/alipay-pool/accounts", headers=headers
    ).json()
    assert sum(float(item["current_balance"]) for item in after_confirm) == 0
    assert sum(float(item["reserved_balance"]) for item in after_confirm) == 0

    reversed_order = client.post(
        f"/api/v1/orders/{order['id']}/status",
        headers=headers,
        json={"status": "REVERSED", "reason": "收银台金额录错"},
    )
    assert reversed_order.status_code == 200, reversed_order.text
    after_reverse = client.get(
        "/api/v1/alipay-pool/accounts", headers=headers
    ).json()
    assert sum(float(item["current_balance"]) for item in after_reverse) == 120
    latest_plan = client.get(
        f"/api/v1/alipay-pool/orders/{order['id']}/plan", headers=headers
    ).json()
    assert latest_plan["status"] == "REVERSED"


def test_soft_cap_plan_uses_remaining_balance_first(client: TestClient):
    _, headers = create_jijihong_context(client)
    device_id = create_device(client, headers)
    create_account(
        client, headers, device_id=device_id, alias="acc-60", balance="60"
    )
    plan = preview_plan(client, headers, "65")
    assert plan["total_balance"] == "60.00"
    assert plan["total_external_cash"] == "5.00"
    assert plan["soft_cap_exceeded"] is False


def test_coupon_is_separate_and_first_day_stops_near_target(client: TestClient):
    _, headers = create_jijihong_context(client)
    device_id = create_device(client, headers)
    coupon_account = client.post(
        "/api/v1/alipay-pool/accounts",
        headers=headers,
        json={
            "device_id": device_id,
            "alias": "coupon-account",
            "current_balance": "45",
            "status": "ACTIVE",
            "phase": "DAY2_ACTIVE",
            "create_coupon": True,
        },
    )
    assert coupon_account.status_code == 201, coupon_account.text
    preview = preview_plan(client, headers, "65")
    created = create_reserved_order(client, headers, amount="65", plan=preview)
    coupon_order = created["order"]
    plan = created["payment_plan"]
    assert plan["total_coupon"] == "20.00"
    assert plan["total_balance"] == "45.00"
    confirmed = client.post(
        f"/api/v1/alipay-pool/payment-plans/{plan['id']}/confirm",
        headers=headers,
        json={},
    )
    assert confirmed.status_code == 200, confirmed.text
    saved_order = client.get(
        f"/api/v1/orders/{coupon_order['id']}", headers=headers
    ).json()
    assert saved_order["coupon_amount"] == "20.00"
    assert saved_order["actual_paid"] == "45.00"

    first_day_account = client.post(
        "/api/v1/alipay-pool/accounts",
        headers=headers,
        json={
            "device_id": device_id,
            "alias": "first-day-account",
            "current_balance": "160",
            "status": "ACTIVE",
            "phase": "NOT_STARTED",
            "birthday_set_date": date.today().isoformat(),
            "create_coupon": True,
        },
    )
    assert first_day_account.status_code == 201, first_day_account.text
    first_day_preview = preview_plan(client, headers, "60")
    first_day_created = create_reserved_order(client, headers, amount="60", plan=first_day_preview)
    first_day_plan = first_day_created["payment_plan"]
    chosen = next(
        item
        for item in first_day_plan["allocations"]
        if item["account_id"] == first_day_account.json()["id"]
    )
    assert chosen["balance_after"] == "100.00"
    client.post(
        f"/api/v1/alipay-pool/payment-plans/{first_day_plan['id']}/confirm",
        headers=headers,
        json={},
    )
    account = next(
        item
        for item in client.get("/api/v1/alipay-pool/accounts", headers=headers).json()
        if item["id"] == first_day_account.json()["id"]
    )
    assert account["phase"] == "WAITING_COUPON"
    assert account["coupon_status"] == "PENDING"


def test_excel_preview_ignores_j_to_o_and_commits_ready_rows(client: TestClient):
    _, headers = create_jijihong_context(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["设备", "类型", "邮箱前三位", "公司", "企业id", "步骤", "造号日期", "设置生日时间", "账户余额", "忽略列"]
    )
    sheet.append(["手机01", "自有", "abc", "公司A", "E-1", "1", None, "8.20", 60, "不得导入"])
    sheet["C2"].fill = PatternFill("solid", fgColor="FFFFFF")
    stream = BytesIO()
    workbook.save(stream)

    preview = client.post(
        "/api/v1/alipay-pool/imports/preview",
        headers=headers,
        files={
            "file": (
                "季季红-充值账号-余额.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    batch = preview.json()
    assert batch["total_rows"] == 1
    assert batch["rows"][0]["parsed_data"]["alias"] == "abc"
    assert "J" not in batch["rows"][0]["raw_data"]

    committed = client.post(
        f"/api/v1/alipay-pool/imports/{batch['id']}/commit",
        headers=headers,
        json={"row_ids": [batch["rows"][0]["id"]]},
    )
    assert committed.status_code == 200, committed.text
    assert committed.json()["imported_rows"] == 1
    accounts = client.get("/api/v1/alipay-pool/accounts", headers=headers).json()
    assert accounts[0]["alias"] == "abc"
    assert accounts[0]["opening_balance"] == "60.00"
    repeated_commit = client.post(
        f"/api/v1/alipay-pool/imports/{batch['id']}/commit",
        headers=headers,
        json={"row_ids": [batch["rows"][0]["id"]]},
    )
    assert repeated_commit.status_code == 409
    repeated_edit = client.patch(
        f"/api/v1/alipay-pool/imports/{batch['id']}/rows/{batch['rows'][0]['id']}",
        headers=headers,
        json={"parsed_data": batch["rows"][0]["parsed_data"]},
    )
    assert repeated_edit.status_code == 409


def test_soft_cap_reason_finance_reports_and_reconciliation(client: TestClient):
    _, headers = create_jijihong_context(client)
    device_id = create_device(client, headers)
    account_id = create_account(
        client, headers, device_id=device_id, alias="cash-gap", balance="60"
    )
    preview = preview_plan(client, headers, "70")
    assert preview["total_external_cash"] == "10.00"
    body = {
        "idempotency_key": "soft-cap-order-0001",
        "business_date": date.today().isoformat(),
        "order_amount": "70",
        "customer_received_amount": "70",
        "allocations": [
            {
                "account_id": item["account_id"],
                "coupon_amount": item["coupon_amount"],
                "balance_amount": item["balance_amount"],
                "external_cash_amount": item["external_cash_amount"],
            }
            for item in preview["allocations"]
        ],
    }
    rejected = client.post(
        "/api/v1/alipay-pool/orders", headers=headers, json=body
    )
    assert rejected.status_code == 422
    body["external_cash_override_reason"] = "客户订单必须完成"
    created = client.post(
        "/api/v1/alipay-pool/orders", headers=headers, json=body
    )
    assert created.status_code == 201, created.text
    repeated = client.post(
        "/api/v1/alipay-pool/orders", headers=headers, json=body
    )
    assert repeated.status_code == 201, repeated.text
    assert repeated.json()["order"]["id"] == created.json()["order"]["id"]
    assert repeated.json()["payment_plan"]["id"] == created.json()["payment_plan"]["id"]
    confirmed = client.post(
        f"/api/v1/alipay-pool/payment-plans/{created.json()['payment_plan']['id']}/confirm",
        headers=headers,
        json={},
    )
    assert confirmed.status_code == 200, confirmed.text

    summary = client.get(
        "/api/v1/reports/jijihong/summary",
        headers=headers,
        params={"date_from": date.today().isoformat(), "date_to": date.today().isoformat()},
    )
    assert summary.status_code == 200, summary.text
    assert summary.json()["customer_received"] == "70.00"
    assert summary.json()["balance_used"] == "60.00"
    assert summary.json()["external_cash"] == "10.00"
    assert summary.json()["cost"] == "70.00"
    assert float(summary.json()["profit"]) == 0
    assert summary.json()["soft_cap_exceeded_count"] == 1
    assert client.get(
        "/api/v1/reports/jijihong/breakdown",
        headers=headers,
        params={
            "group_by": "account",
            "date_from": date.today().isoformat(),
            "date_to": date.today().isoformat(),
        },
    ).json()[0]["entity_id"] == account_id
    exported = client.get(
        "/api/v1/reports/jijihong/orders/export",
        headers=headers,
        params={"export_format": "xlsx"},
    )
    assert exported.status_code == 200
    assert exported.content.startswith(b"PK")

    stock_id = create_account(
        client, headers, device_id=device_id, alias="stocktake", balance="100"
    )
    draft = client.post(
        "/api/v1/alipay-pool/reconciliations",
        headers=headers,
        json={
            "business_date": date.today().isoformat(),
            "items": [{"account_id": stock_id, "actual_balance": "95", "reason": "收银台盘点"}],
        },
    )
    assert draft.status_code == 201, draft.text
    adjusted = client.post(
        f"/api/v1/alipay-pool/accounts/{stock_id}/adjust",
        headers=headers,
        json={"amount": "-1", "reason": "测试并发变化"},
    )
    assert adjusted.status_code == 200, adjusted.text
    stale_confirm = client.post(
        f"/api/v1/alipay-pool/reconciliations/{draft.json()['id']}/confirm",
        headers=headers,
    )
    assert stale_confirm.status_code == 409

    fresh = client.post(
        "/api/v1/alipay-pool/reconciliations",
        headers=headers,
        json={
            "business_date": date.today().isoformat(),
            "items": [{"account_id": stock_id, "actual_balance": "95", "reason": "收银台盘点"}],
        },
    )
    confirmed_reconciliation = client.post(
        f"/api/v1/alipay-pool/reconciliations/{fresh.json()['id']}/confirm",
        headers=headers,
    )
    assert confirmed_reconciliation.status_code == 200, confirmed_reconciliation.text
    assert confirmed_reconciliation.json()["items"][0]["difference"] == "-4.00"
    reversed_reconciliation = client.post(
        f"/api/v1/alipay-pool/reconciliations/{fresh.json()['id']}/reverse",
        headers=headers,
    )
    assert reversed_reconciliation.status_code == 200, reversed_reconciliation.text
    restored = next(
        item
        for item in client.get("/api/v1/alipay-pool/accounts", headers=headers).json()
        if item["id"] == stock_id
    )
    assert restored["current_balance"] == "99.00"


def test_device_has_at_most_five_non_exhausted_accounts(client: TestClient):
    _, headers = create_jijihong_context(client)
    device_id = create_device(client, headers)
    for index in range(5):
        create_account(
            client,
            headers,
            device_id=device_id,
            alias=f"slot-{index}",
            balance="1",
        )
    overflow = client.post(
        "/api/v1/alipay-pool/accounts",
        headers=headers,
        json={
            "device_id": device_id,
            "alias": "slot-overflow",
            "current_balance": "1",
            "status": "ACTIVE",
            "phase": "DAY2_ACTIVE",
            "create_coupon": False,
        },
    )
    assert overflow.status_code == 409
    exhausted = client.post(
        "/api/v1/alipay-pool/accounts",
        headers=headers,
        json={
            "device_id": device_id,
            "alias": "archived-empty",
            "current_balance": "0",
            "status": "EXHAUSTED",
            "phase": "EXHAUSTED",
            "create_coupon": False,
        },
    )
    assert exhausted.status_code == 201, exhausted.text


def test_competing_orders_cancel_release_and_replace_reservation(client: TestClient):
    _, headers = create_jijihong_context(client)
    device_id = create_device(client, headers)
    first_id = create_account(
        client, headers, device_id=device_id, alias="reserve-first", balance="60"
    )
    stale_preview = preview_plan(client, headers, "60")
    first = create_reserved_order(
        client, headers, amount="60", plan=stale_preview
    )
    competing = client.post(
        "/api/v1/alipay-pool/orders",
        headers=headers,
        json={
            "business_date": date.today().isoformat(),
            "order_amount": "60",
            "customer_received_amount": "60",
            "allocations": [
                {
                    "account_id": item["account_id"],
                    "coupon_amount": item["coupon_amount"],
                    "balance_amount": item["balance_amount"],
                    "external_cash_amount": item["external_cash_amount"],
                }
                for item in stale_preview["allocations"]
            ],
        },
    )
    assert competing.status_code == 409
    reserved = next(
        item
        for item in client.get("/api/v1/alipay-pool/accounts", headers=headers).json()
        if item["id"] == first_id
    )
    assert reserved["reserved_balance"] == "60.00"

    cancelled = client.post(
        f"/api/v1/orders/{first['order']['id']}/status",
        headers=headers,
        json={"status": "CANCELLED", "reason": "客户取消"},
    )
    assert cancelled.status_code == 200, cancelled.text
    released = next(
        item
        for item in client.get("/api/v1/alipay-pool/accounts", headers=headers).json()
        if item["id"] == first_id
    )
    assert released["reserved_balance"] == "0.00"

    second_id = create_account(
        client, headers, device_id=device_id, alias="reserve-second", balance="60"
    )
    draft = create_reserved_order(client, headers, amount="60")
    old_account_id = draft["payment_plan"]["allocations"][0]["account_id"]
    replacement_id = second_id if old_account_id == first_id else first_id
    replaced = client.put(
        f"/api/v1/alipay-pool/orders/{draft['order']['id']}/reservation",
        headers=headers,
        json={
            "order_amount": "60",
            "customer_received_amount": "58",
            "allocations": [
                {
                    "account_id": replacement_id,
                    "coupon_amount": "0",
                    "balance_amount": "60",
                    "external_cash_amount": "0",
                }
            ],
            "note": "手工替换账号",
        },
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["payment_plan"]["allocations"][0]["account_id"] == replacement_id
    assert replaced.json()["order"]["customer_received_amount"] == "58.00"
    changed_at_confirm = client.post(
        f"/api/v1/alipay-pool/payment-plans/{replaced.json()['payment_plan']['id']}/confirm",
        headers=headers,
        json={
            "allocations": [
                {
                    "account_id": replacement_id,
                    "coupon_amount": "0",
                    "balance_amount": "59",
                    "external_cash_amount": "1",
                }
            ]
        },
    )
    assert changed_at_confirm.status_code == 409
