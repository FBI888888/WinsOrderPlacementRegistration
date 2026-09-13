import hashlib
import re
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.alipay_pool.models import (
    AccountPhase,
    AccountStatus,
    AlipayImportBatch,
    AlipayImportRow,
    ImportBatchStatus,
    ImportRowStatus,
)


def _text(value) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _balance(value) -> tuple[Decimal | None, bool, str | None]:
    if value is None:
        return None, False, "缺少账户余额"
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01")), False, None
    raw = str(value).strip().lower()
    coupon_marker = "juan" in raw or "卷" in raw or "券" in raw
    match = re.search(r"-?\d+(?:\.\d+)?", raw)
    if not match:
        return None, coupon_marker, f"无法识别余额：{value}"
    try:
        return Decimal(match.group()).quantize(Decimal("0.01")), coupon_marker, None
    except InvalidOperation:
        return None, coupon_marker, f"无法识别余额：{value}"


def _birthday(value) -> tuple[date | None, str | None]:
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    raw = str(value).strip().replace("月", ".").replace("日", "")
    match = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})", raw)
    if not match:
        return None, f"无法识别生日设置日期：{value}"
    try:
        return date(date.today().year, int(match.group(1)), int(match.group(2))), None
    except ValueError:
        return None, f"无法识别生日设置日期：{value}"


def _fill_rgb(cell) -> str | None:
    color = cell.fill.fgColor
    if color.type != "rgb" or not color.rgb:
        return None
    return color.rgb[-6:].upper()


def _phase_from_fill(rgb: str | None) -> tuple[str, str, bool, list[str]]:
    if rgb == "00B050":
        return (
            AccountPhase.WAITING_COUPON.value,
            AccountStatus.ACTIVE.value,
            False,
            [],
        )
    if rgb == "FFC000":
        return (
            AccountPhase.DAY2_ACTIVE.value,
            AccountStatus.ACTIVE.value,
            True,
            [],
        )
    if rgb == "FF0000":
        return (
            AccountPhase.EXHAUSTED.value,
            AccountStatus.EXHAUSTED.value,
            False,
            [],
        )
    if rgb in (None, "FFFFFF", "000000"):
        return (
            AccountPhase.NOT_STARTED.value,
            AccountStatus.ACTIVE.value,
            False,
            [],
        )
    return (
        AccountPhase.NOT_STARTED.value,
        AccountStatus.ACTIVE.value,
        False,
        [f"未识别的账号底色 {rgb}，请确认阶段"],
    )


def create_preview(
    db: Session,
    *,
    tenant_id: int,
    user_id: int,
    filename: str,
    content: bytes,
) -> AlipayImportBatch:
    file_hash = hashlib.sha256(content).hexdigest()
    existing = db.scalar(
        select(AlipayImportBatch).where(
            AlipayImportBatch.tenant_id == tenant_id,
            AlipayImportBatch.file_hash == file_hash,
        )
    )
    if existing:
        return existing
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取 Excel 文件") from exc
    sheet = workbook[workbook.sheetnames[0]]
    batch = AlipayImportBatch(
        tenant_id=tenant_id,
        filename=filename[:255],
        file_hash=file_hash,
        sheet_name=sheet.title[:100],
        status=ImportBatchStatus.PREVIEW.value,
        created_by=user_id,
    )
    db.add(batch)
    db.flush()

    last_device: str | None = None
    parsed_rows: list[AlipayImportRow] = []
    aliases: dict[str, list[int]] = {}
    for row_number in range(2, sheet.max_row + 1):
        cells = [sheet.cell(row_number, column) for column in range(1, 10)]
        values = [cell.value for cell in cells]
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        if _text(values[0]):
            last_device = _text(values[0])
        alias = _text(values[2])
        # 说明行、仅用于设备分组或保存企业元数据的行，不是支付宝账号。
        if not alias and _text(values[8]) is None:
            continue
        balance, coupon_marker, balance_error = _balance(values[8])
        birthday, birthday_error = _birthday(values[7])
        phase, account_status, color_coupon, warnings = _phase_from_fill(_fill_rgb(cells[2]))
        errors = [error for error in (balance_error, birthday_error) if error]
        if not last_device:
            errors.append("缺少所属设备")
        if not alias:
            errors.append("缺少支付宝账号别名")
        if balance is not None and balance < 0:
            errors.append("账户余额不能为负数")
        coupon_available = coupon_marker or color_coupon
        if coupon_marker:
            warnings.append("余额文本包含优惠券标记，已按20元可用券解析")
        if alias:
            aliases.setdefault(alias.lower(), []).append(row_number)
        parsed = {
            "device_name": last_device,
            "account_category": _text(values[1]),
            "alias": alias,
            "company": _text(values[3]),
            "enterprise_id": _text(values[4]),
            "step": _text(values[5]),
            "created_date": _text(values[6]),
            "birthday_set_date": birthday.isoformat() if birthday else None,
            "current_balance": str(balance) if balance is not None else None,
            "phase": phase,
            "status": account_status,
            "coupon_available": coupon_available,
        }
        row_status = (
            ImportRowStatus.REVIEW.value
            if errors or warnings
            else ImportRowStatus.READY.value
        )
        row = AlipayImportRow(
            tenant_id=tenant_id,
            batch_id=batch.id,
            row_number=row_number,
            raw_data={chr(65 + index): _text(value) for index, value in enumerate(values)},
            parsed_data=parsed,
            warnings=warnings,
            errors=errors,
            status=row_status,
        )
        db.add(row)
        parsed_rows.append(row)

    duplicate_rows = {
        row_number
        for row_numbers in aliases.values()
        if len(row_numbers) > 1
        for row_number in row_numbers
    }
    for row in parsed_rows:
        if row.row_number in duplicate_rows:
            row.warnings = [*(row.warnings or []), "账号别名重复；系统允许重复并使用内部ID区分"]
            if row.status == ImportRowStatus.READY.value:
                row.status = ImportRowStatus.REVIEW.value

    batch.total_rows = len(parsed_rows)
    batch.ready_rows = sum(row.status == ImportRowStatus.READY.value for row in parsed_rows)
    batch.review_rows = sum(row.status == ImportRowStatus.REVIEW.value for row in parsed_rows)
    db.commit()
    db.refresh(batch)
    return batch


def coupon_available_at(birthday_set_date: date | None) -> datetime | None:
    if not birthday_set_date:
        return None
    return datetime.combine(
        birthday_set_date + timedelta(days=1),
        time.min,
        tzinfo=timezone.utc,
    )
